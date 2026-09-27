"""Build realistic timetable files in every format the upload accepts.

Dates are generated relative to today so the suite never goes stale and every
reminder lead time is still in the future.
"""
import io
from datetime import date, timedelta

# (course code, days from today, column) — column 0 = 9:00am, 1 = 1:00pm
ROWS = [
    ("CSC201", 10, 0), ("MTH211", 10, 1),
    ("PHY204", 11, 0), ("CSC205", 11, 1),
    ("ECO101", 12, 0), ("CSC209", 12, 1),
    ("SENG101", 13, 0), ("ENG101", 13, 1),
    ("STA202", 14, 0), ("CSC211", 14, 1),
]
TIMES = ["09:00", "13:00"]


def day(offset: int) -> date:
    return date.today() + timedelta(days=offset)


def human(d: date) -> str:
    return f"{d.strftime('%A')} {d.day} {d.strftime('%B %Y')}"


def grid() -> list[list[str]]:
    """Rows of [date label, 9am course, 1pm course]."""
    by_day: dict[int, list[str]] = {}
    for code, offset, col in ROWS:
        by_day.setdefault(offset, ["", ""])[col] = code
    return [[human(day(o)), *cells] for o, cells in sorted(by_day.items())]


def expected(code: str) -> tuple[str, str]:
    """(ISO date, HH:MM) the parser should produce for a course."""
    for c, offset, col in ROWS:
        if c == code:
            return day(offset).isoformat(), TIMES[col]
    raise KeyError(code)


HEADER = ["DATE", "9:00am", "1:00pm"]


def as_txt() -> bytes:
    lines = ["FACULTY OF SCIENCE — SECOND SEMESTER EXAMINATION TIMETABLE", " | ".join(HEADER)]
    lines += [" | ".join(r) for r in grid()]
    return "\n".join(lines).encode()


def as_csv() -> bytes:
    lines = [",".join(HEADER)] + [",".join(r) for r in grid()]
    return "\n".join(lines).encode()


def as_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for r in grid():
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def as_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("FACULTY OF SCIENCE — EXAMINATION TIMETABLE")
    table = doc.add_table(rows=1, cols=3)
    for cell, h in zip(table.rows[0].cells, HEADER):
        cell.text = h
    for r in grid():
        cells = table.add_row().cells
        for cell, v in zip(cells, r):
            cell.text = v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def as_pdf() -> bytes:
    """A minimal, valid, text-based PDF (no extra dependency needed)."""
    lines = ["EXAMINATION TIMETABLE", "   ".join(HEADER)] + ["   ".join(r) for r in grid()]
    stream = "BT /F1 10 Tf 40 780 Td 14 TL\n"
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream += f"({safe}) Tj T*\n"
    stream += "ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def as_png() -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 400), "white")
    draw = ImageDraw.Draw(img)
    y = 10
    for line in [" | ".join(HEADER)] + [" | ".join(r) for r in grid()]:
        draw.text((10, y), line, fill="black")
        y += 24
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def big_txt(n_days: int = 400) -> bytes:
    """A very large timetable: hundreds of days x 2 slots of filler courses."""
    lines = [" | ".join(HEADER)]
    for i in range(n_days):
        lines.append(f"{human(day(10 + i))} | ZZZ{1000 + i} | YYY{1000 + i}")
    return "\n".join(lines).encode()


FORMATS = {
    "txt": as_txt,
    "csv": as_csv,
    "xlsx": as_xlsx,
    "docx": as_docx,
    "pdf": as_pdf,
}
